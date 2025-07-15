#!/usr/bin/env python3
"""
Analyze the Excel document structure and answer patterns
"""

import sys
from pathlib import Path
import openpyxl

# Add the backend directory to Python path
backend_dir = Path(__file__).parent
sys.path.insert(0, str(backend_dir))


def analyze_excel_structure():
    """Analyze the complete Excel document structure and answer patterns"""
    # Use one of the uploaded files
    esg_file = Path("/Users/axerroce/esg-checklist-ai/backend/uploads/2025/07/3_esg-doc_76a28472_ESG-Internal_Audit-Checklist-_Revised_Dummy_Data.4.xlsx")
    
    if not esg_file.exists():
        print(f"ESG file not found: {esg_file}")
        return
    
    print(f"Analyzing Excel structure: {esg_file.name}")
    
    wb = openpyxl.load_workbook(esg_file)
    
    # Show all worksheets
    print(f"\nWorksheets in file:")
    for i, ws in enumerate(wb.worksheets):
        print(f"  {i+1}. '{ws.title}' ({ws.max_row} rows, {ws.max_column} columns)")
    
    # Find the ESG questionnaires sheet
    target_sheet = None
    for ws in wb.worksheets:
        if "esg" in ws.title.lower() and (
            "questionnaire" in ws.title.lower() or "question" in ws.title.lower()
        ):
            target_sheet = ws
            break
    
    if not target_sheet:
        print("No ESG questionnaires sheet found!")
        return
    
    print(f"\nAnalyzing sheet: '{target_sheet.title}'")
    print(f"Dimensions: {target_sheet.max_row} rows × {target_sheet.max_column} columns")
    
    # Analyze the header row
    print(f"\nHeader row analysis:")
    header_row = 1
    headers = []
    for col in range(1, target_sheet.max_column + 1):
        cell_value = target_sheet.cell(row=header_row, column=col).value
        headers.append(cell_value)
        print(f"  Column {col}: '{cell_value}'")
    
    # Show some sample data rows
    print(f"\nSample data rows (first 10 non-empty rows):")
    sample_count = 0
    for row in range(2, target_sheet.max_row + 1):
        row_data = []
        has_data = False
        for col in range(1, target_sheet.max_column + 1):
            cell_value = target_sheet.cell(row=row, column=col).value
            row_data.append(str(cell_value) if cell_value else "")
            if cell_value and str(cell_value).strip():
                has_data = True
        
        if has_data:
            sample_count += 1
            print(f"\nRow {row}:")
            for i, (header, value) in enumerate(zip(headers, row_data)):
                if value.strip():  # Only show non-empty values
                    print(f"  {header}: '{value}'")
            
            if sample_count >= 10:
                break
    
    # Analyze answer patterns
    print(f"\n" + "="*50)
    print("ANSWER PATTERN ANALYSIS")
    print("="*50)
    
    # Look for assessment/answer columns
    assessment_columns = []
    for i, header in enumerate(headers):
        if header and ("assessment" in str(header).lower() or 
                      "answer" in str(header).lower() or
                      "response" in str(header).lower() or
                      "provide" in str(header).lower()):
            assessment_columns.append((i + 1, header))  # Column numbers are 1-based
    
    print(f"\nPotential answer/assessment columns found:")
    for col_num, header in assessment_columns:
        print(f"  Column {col_num}: '{header}'")
    
    # Analyze answer patterns for each assessment column
    for col_num, header in assessment_columns:
        print(f"\nAnalyzing answers in column '{header}' (Column {col_num}):")
        
        answer_patterns = {}
        non_empty_count = 0
        
        for row in range(2, target_sheet.max_row + 1):
            cell_value = target_sheet.cell(row=row, column=col_num).value
            if cell_value and str(cell_value).strip():
                answer = str(cell_value).strip()
                non_empty_count += 1
                
                # Categorize the answer
                answer_lower = answer.lower()
                if answer_lower in ["yes", "y", "true", "✓", "✔"]:
                    category = "Complete/Yes"
                elif answer_lower in ["no", "n", "false", "✗", "✘"]:
                    category = "Incomplete/No"
                elif answer_lower in ["n/a", "na", "not applicable", "not available"]:
                    category = "Not Applicable/Available"
                elif len(answer) > 20:  # Detailed text responses
                    category = "Detailed Response"
                else:
                    category = "Other"
                
                if category not in answer_patterns:
                    answer_patterns[category] = []
                
                answer_patterns[category].append(answer)
        
        print(f"  Total non-empty answers: {non_empty_count}")
        print(f"  Answer patterns found:")
        
        for category, answers in answer_patterns.items():
            unique_answers = list(set(answers))[:5]  # Show up to 5 unique examples
            print(f"    {category} ({len(answers)} answers):")
            for answer in unique_answers:
                print(f"      - '{answer}'")
            if len(set(answers)) > 5:
                print(f"      ... and {len(set(answers)) - 5} more unique answers")
    
    # Check for reference codes and their patterns
    print(f"\n" + "="*50)
    print("REFERENCE CODE ANALYSIS")
    print("="*50)
    
    reference_patterns = {
        "Environmental": 0,
        "Social": 0,
        "Governance": 0,
        "Follow-up questions": 0
    }
    
    for row in range(2, target_sheet.max_row + 1):
        for col in range(1, target_sheet.max_column + 1):
            cell_value = target_sheet.cell(row=row, column=col).value
            if cell_value and "ESG-" in str(cell_value):
                ref_code = str(cell_value)
                
                if "Environment" in ref_code:
                    reference_patterns["Environmental"] += 1
                elif "Social" in ref_code:
                    reference_patterns["Social"] += 1
                elif "Governance" in ref_code:
                    reference_patterns["Governance"] += 1
                
                # Check for follow-up questions (ending with b, c, etc.)
                import re
                if re.search(r'[b-z]:', ref_code, re.IGNORECASE):
                    reference_patterns["Follow-up questions"] += 1
    
    print("Reference code distribution:")
    for category, count in reference_patterns.items():
        print(f"  {category}: {count}")


if __name__ == "__main__":
    analyze_excel_structure()